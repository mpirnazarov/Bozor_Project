"""Do'kon va infra ro'yxati importi (bitta Excel faylda).

Fayl ko'rinishi (Chorsu "ATROF" namunasi):

  № | Tadbirkorlar ro`yxati | JSHSHIR | Telefon | Kv.metr | Infra stavkasi |
  Infra summasi | kv.metr | Ijara stavkasi | Ijara summasi | Komunal | ...

Faylda infra ustunlari bo'lmasligi ham mumkin (masalan Chorsu "Milliy
kiyimlar" — faqat ijara). Unda hamma qator oddiy do'kon bo'ladi.

Magazin raqami: agar faylda alohida "Do'kon raqami" ustuni bo'lsa — o'sha,
aks holda "№" ustuni ishlatiladi (Milliy kiyimlarda ular farq qiladi:
№ 96 -> do'kon 98).

Qoidalar:
  * `Infra summasi` to'ldirilgan  -> INFRA do'kon (`infra_shops`)
  * `Ijara summasi` to'ldirilgan  -> oddiy do'kon (`shops`)
  * Ega raqami XONALAR SONI bo'yicha ajratiladi:
      9 xonali  -> INN      (yuridik shaxs)
      14 xonali -> JSHSHIR  (jismoniy shaxs)
  * `Jami to'lov`, `Komunal`, `Poteriya` — O'QILMAYDI (bazada maydon yo'q
    va `Jami to'lov` da xatolik bor).

`preview=True` bo'lsa bazaga hech narsa yozilmaydi — faqat nima
o'zgarishi hisoblanadi.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Counterparty, InfraShop, Shop


class StructureError(Exception):
    """Excel strukturasi kutilganidan farq qilganda."""


def _norm(v) -> str:
    return "".join(str(v or "").lower().split()).replace("`", "'").replace("’", "'")


def _dec(v) -> Decimal:
    if v is None:
        return Decimal(0)
    if isinstance(v, (int, float)):
        try:
            return Decimal(str(round(float(v), 2)))
        except (InvalidOperation, ValueError):
            return Decimal(0)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not s:
        return Decimal(0)
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        return Decimal(0)


def _digits(v) -> str:
    return "".join(ch for ch in str(v or "") if ch.isdigit())


def id_kind(num: str) -> str | None:
    """Raqam turini xonalar soniga qarab aniqlaydi."""
    if len(num) == 9:
        return "inn"
    if len(num) == 14:
        return "jshshir"
    return None


@dataclass
class Row:
    no: str
    shop_no: str   # "Do'kon raqami" ustuni; bo'lmasa `no` bilan bir xil
    name: str
    ident: str
    id_type: str | None
    phone: str
    infra_area: Decimal
    infra_sum: Decimal
    rent_area: Decimal
    rent_sum: Decimal
    kind: str  # "infra" | "shop"


@dataclass
class ShopsInfraResult:
    rows_read: int = 0
    shops_new: int = 0
    shops_updated: int = 0
    infra_new: int = 0
    infra_updated: int = 0
    cp_new: int = 0
    cp_updated: int = 0
    inn_count: int = 0
    jshshir_count: int = 0
    no_id: int = 0
    bad_id: list[str] = field(default_factory=list)
    shops_total: Decimal = Decimal(0)
    infra_total: Decimal = Decimal(0)
    skipped: list[dict] = field(default_factory=list)
    detected_columns: dict = field(default_factory=dict)
    sample: list[dict] = field(default_factory=list)


def _find_columns(rows: list) -> tuple[int, dict[str, int]]:
    """Sarlavha qatorini va ustun indekslarini topadi.

    Faylda IKKITA "kv.metr" ustuni bor (biri infra, biri ijara) — ular
    faqat harf registri bilan farq qiladi, shuning uchun maydon
    ustunlari "stavka" ustunidan CHAPDAGI ustun sifatida aniqlanadi.
    """
    for i, r in enumerate(rows[:12]):
        cells = [_norm(c) for c in r]
        col: dict[str, int] = {}
        for j, c in enumerate(cells):
            if not c:
                continue
            if "infra" in c and "summa" in c:
                col["infra_sum"] = j
            elif "infra" in c and "stavka" in c:
                col["infra_rate"] = j
            elif "ijara" in c and "summa" in c:
                col["rent_sum"] = j
            elif "ijara" in c and "stavka" in c:
                col["rent_rate"] = j
            elif "jshshir" in c or c in ("inn", "инн", "stir"):
                col["ident"] = j
            elif "tadbirkor" in c or "ro'yxat" in c or "kontragent" in c:
                col["name"] = j
            elif "telefon" in c:
                col["phone"] = j
            elif ("дукон" in c or "do'kon" in c or "dukon" in c) and (
                "рак" in c or "raqam" in c or "nomer" in c
            ):
                col["shop_no"] = j
            elif c in ("№", "n", "no", "nomer"):
                col["no"] = j
        # Infra ustunlari majburiy emas — faqat ijarali fayl ham bo'ladi
        if ("infra_sum" in col or "rent_sum" in col) and "ident" in col:
            if "infra_rate" in col and col["infra_rate"] > 0:
                col["infra_area"] = col["infra_rate"] - 1
            if "rent_rate" in col and col["rent_rate"] > 0:
                col["rent_area"] = col["rent_rate"] - 1
            return i, col
    raise StructureError(
        "Excel strukturasi mos kelmadi. Kerakli ustunlar: "
        "Tadbirkorlar ro'yxati, JSHSHIR va Ijara summasi (yoki Infra summasi)."
    )


def parse_file(content: bytes) -> tuple[list[Row], dict]:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise StructureError(f"Faylni ochib bo'lmadi (.xlsx kerak): {exc}") from exc

    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        raise StructureError("Fayl bo'sh")
    hdr, col = _find_columns(rows)

    def cell(r, key):
        j = col.get(key)
        return r[j] if (j is not None and j < len(r)) else None

    out: list[Row] = []
    for r in rows[hdr + 1:]:
        r = list(r)
        no = str(cell(r, "no") or "").strip()
        shop_no = str(cell(r, "shop_no") or "").strip() or no
        name = str(cell(r, "name") or "").strip()
        if not no and not name:
            continue
        infra_sum = _dec(cell(r, "infra_sum"))
        rent_sum = _dec(cell(r, "rent_sum"))
        if infra_sum <= 0 and rent_sum <= 0:
            continue  # summasi yo'q — o'tkazamiz
        ident = _digits(cell(r, "ident"))
        out.append(Row(
            no=no, shop_no=shop_no, name=name, ident=ident, id_type=id_kind(ident),
            phone=str(cell(r, "phone") or "").strip(),
            infra_area=_dec(cell(r, "infra_area")), infra_sum=infra_sum,
            rent_area=_dec(cell(r, "rent_area")), rent_sum=rent_sum,
            kind="infra" if infra_sum > 0 else "shop",
        ))
    labels = {"no": "№", "shop_no": "Do'kon raqami", "name": "Tadbirkor", "ident": "JSHSHIR/INN", "phone": "Telefon",
              "infra_area": "Infra kv.m", "infra_rate": "Infra stavkasi",
              "infra_sum": "Infra summasi", "rent_area": "Ijara kv.m",
              "rent_rate": "Ijara stavkasi", "rent_sum": "Ijara summasi"}
    return out, {labels.get(k, k): v for k, v in col.items()}


async def import_shops_infra(
    db: AsyncSession,
    content: bytes,
    market_id: int,
    prefix: str = "",
    preview: bool = True,
) -> ShopsInfraResult:
    parsed, cols = parse_file(content)
    res = ShopsInfraResult(rows_read=len(parsed), detected_columns=cols)
    if not parsed:
        raise StructureError("Faylda summasi bor birorta qator topilmadi")

    shop_rows = [r for r in parsed if r.kind == "shop"]
    infra_rows = [r for r in parsed if r.kind == "infra"]

    shop_ids = [f"{prefix}{r.shop_no}" for r in shop_rows]
    existing_shops = {
        s.shop_id: s for s in (await db.execute(
            select(Shop).where(Shop.market_id == market_id, Shop.shop_id.in_(shop_ids))
        )).scalars()
    } if shop_ids else {}

    infra_names = [r.name for r in infra_rows]
    existing_infra = {
        i.name: i for i in (await db.execute(
            select(InfraShop).where(InfraShop.market_id == market_id,
                                    InfraShop.name.in_(infra_names))
        )).scalars()
    } if infra_names else {}

    idents = {r.ident for r in parsed if r.ident}
    existing_cp = {
        c.inn: c for c in (await db.execute(
            select(Counterparty).where(Counterparty.inn.in_(idents))
        )).scalars()
    } if idents else {}

    for r in parsed:
        if not r.ident:
            res.no_id += 1
            res.skipped.append({"no": r.shop_no, "name": r.name, "reason": "ID yo'q"})
        elif r.id_type is None:
            res.bad_id.append(f"{r.no}: {r.ident} ({len(r.ident)} xonali)")
        elif r.id_type == "inn":
            res.inn_count += 1
        else:
            res.jshshir_count += 1

    for r in shop_rows:
        if f"{prefix}{r.shop_no}" in existing_shops:
            res.shops_updated += 1
        res.shops_total += r.rent_sum
    for r in infra_rows:
        if r.name in existing_infra:
            res.infra_updated += 1
        else:
            res.infra_new += 1
        res.infra_total += r.infra_sum
    res.shops_new = len(shop_rows) - res.shops_updated

    new_idents = {r.ident for r in parsed if r.ident and r.ident not in existing_cp}
    res.cp_new = len(new_idents)
    res.cp_updated = len({r.ident for r in parsed if r.ident and r.ident in existing_cp})

    res.sample = [
        {"no": r.shop_no, "shop_id": (f"{prefix}{r.shop_no}" if r.kind == "shop" else None),
         "name": r.name, "ident": r.ident, "id_type": r.id_type or "—",
         "kind": r.kind, "summa": float(r.infra_sum if r.kind == "infra" else r.rent_sum)}
        for r in parsed[:15]
    ]

    if preview:
        return res

    # ===== YOZISH =====
    # Kontragentlar AVVAL (shops.inn / infra_shops.inn FK uchun)
    for r in parsed:
        if not r.ident:
            continue
        cp = existing_cp.get(r.ident)
        if cp is None:
            cp = Counterparty(
                inn=r.ident, name=r.name or f"ID {r.ident}",
                id_type=r.id_type or "inn", phone=r.phone or None,
            )
            db.add(cp)
            existing_cp[r.ident] = cp
        else:
            if r.name and cp.name != r.name:
                cp.name = r.name
            if r.phone and not cp.phone:
                cp.phone = r.phone
            if r.id_type and cp.id_type != r.id_type:
                cp.id_type = r.id_type
    await db.flush()

    for r in shop_rows:
        sid = f"{prefix}{r.shop_no}"
        s = existing_shops.get(sid)
        if s is None:
            db.add(Shop(
                shop_id=sid, market_id=market_id, inn=r.ident or None,
                monthly_rent=r.rent_sum,
                area=float(r.rent_area) if r.rent_area > 0 else None,
                is_active=True, source_sheet="shops-infra-excel",
            ))
        else:
            s.inn = r.ident or s.inn
            s.monthly_rent = r.rent_sum
            if r.rent_area > 0:
                s.area = float(r.rent_area)
            s.source_sheet = "shops-infra-excel"

    for r in infra_rows:
        i = existing_infra.get(r.name)
        if i is None:
            db.add(InfraShop(
                market_id=market_id, name=r.name, inn=r.ident or None,
                monthly_rent=r.infra_sum, is_active=True,
            ))
        else:
            i.inn = r.ident or i.inn
            i.monthly_rent = r.infra_sum
    await db.flush()
    return res
